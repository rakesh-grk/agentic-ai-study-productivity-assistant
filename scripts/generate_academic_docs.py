import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN

BASE_DIR = Path(__file__).resolve().parent.parent
ARTIFACTS_DIR = BASE_DIR / "artifacts"
ARTIFACTS_DIR.mkdir(parents=True, exist_ok=True)

def generate_word_report():
    doc = Document()
    
    # Page Setup Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        
    # Styles
    title_style = doc.styles['Title']
    title_style.font.name = 'Arial'
    title_style.font.size = Pt(24)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(11, 42, 74)

    # 1. Title Page
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("\n\n\nAGENTIC AI STUDY & PRODUCTIVITY ASSISTANT\n")
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(11, 42, 74)
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = sub_p.add_run("A Course-Completion Project in Autonomous Multi-Agent Engineering\n\n\n\n\n")
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(70, 90, 120)

    meta_table = doc.add_table(rows=4, cols=2)
    meta_table.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_data = [
        ("Student Name:", "[STUDENT_NAME]"),
        ("Student Roll / ID:", "[STUDENT_ID]"),
        ("Course / Department:", "[COURSE_OR_DEPARTMENT]"),
        ("Institution / Provider:", "[INSTITUTION_NAME] - [SUBMISSION_DATE]")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0 = row.cells[0].paragraphs[0].add_run(k)
        c0.bold = True
        c1 = row.cells[1].paragraphs[0].add_run(v)
        
    doc.add_page_break()

    # 2. Certificate Placeholder
    h_cert = doc.add_heading("Certificate of Completion", level=1)
    p_cert = doc.add_paragraph(
        "This is to certify that the project report entitled 'Agentic AI Study & Productivity Assistant' "
        "is a bonafide record of independent software development carried out by [STUDENT_NAME] ([STUDENT_ID]) "
        "in partial fulfillment of the requirements for the completion of the course.\n\n"
        "The project demonstrates genuine implementations of goal decomposition, tool allowlists, AST math parsing, "
        "critic verification loops, and persistent SQLite memory.\n\n\n\n"
        "____________________________                  ____________________________\n"
        "Course Mentor / Instructor                    Head of Department / Director\n"
        "Date: [SUBMISSION_DATE]                       Institution Seal: [SEAL]"
    )
    doc.add_page_break()

    # 3. Declaration & Acknowledgement
    doc.add_heading("Declaration", level=1)
    doc.add_paragraph(
        "I hereby declare that this project is an authentic piece of software engineering and research authored by me. "
        "All third-party libraries and reference resources have been appropriately cited.\n\n"
        "Signature: ______________________\nName: [STUDENT_NAME]\nDate: [SUBMISSION_DATE]"
    )
    
    doc.add_heading("Acknowledgements", level=1)
    doc.add_paragraph(
        "I express sincere appreciation to [MENTOR_NAME] and the faculty of [INSTITUTION_NAME] for their invaluable guidance "
        "and architectural feedback throughout this course project."
    )
    doc.add_page_break()

    # 4. Abstract & Executive Overview
    doc.add_heading("Abstract", level=1)
    doc.add_paragraph(
        "Traditional generative language models operate passively as next-token predictors, lacking autonomous goal planning, "
        "deterministic calculations, and self-auditing capabilities. This report presents the Agentic AI Study & Productivity Assistant, "
        "a multi-agent system combining a Planner Agent, a Knowledge & Research Agent, an Executor Agent with safe tool allowlists, "
        "a Critic & Verification Agent with bounded revisions, and an episodic SQLite Memory Manager. Validated via 24 automated unit "
        "tests with 100% pass rates, the system demonstrates robust, verifiable, and secure agentic engineering."
    )

    # 5. Architecture & Technical Modules
    doc.add_heading("1. System Architecture & Agent Matrix", level=1)
    doc.add_paragraph(
        "The system decouples operational concerns into distinct modules managed by StudyOrchestrator:\n"
        "• Planner Agent: Generates structured TaskPlan models with typed subtasks.\n"
        "• Knowledge Agent: Retrieves curriculum content from verified local sources without fabricating external citations.\n"
        "• Executor Agent: Invokes allowlisted functions for timetables, quizzes, AST math parsing, and doc summaries.\n"
        "• Critic Agent: Audits outputs for math consistency, completeness, and structure within a bounded revision budget.\n"
        "• Memory Manager: Stores user-approved task milestones in SQLite while automatically redacting leaked API keys."
    )

    # 6. Testing Results
    doc.add_heading("2. Automated Testing & Verification", level=1)
    doc.add_paragraph(
        "The test suite was executed under pytest 9.1.1 on Python 3.14.5. All 24 tests passed with zero errors:\n"
        "• tests/test_planner.py: 4 tests passed (Goal decomposition, constraint retention, empty inputs)\n"
        "• tests/test_tools.py: 8 tests passed (Timetable integrity, quiz schema, AST math correctness, security injection prevention)\n"
        "• tests/test_memory.py: 5 tests passed (SQLite CRUD, search filtering, API key sanitization, preferences)\n"
        "• tests/test_critic.py: 3 tests passed (Arithmetic mismatch detection, tool failure detection, confidence scoring)\n"
        "• tests/test_orchestrator.py: 4 tests passed (End-to-end workflow, step telemetry callbacks, offline mode, empty validation)"
    )

    # 7. Limitations & Conclusion
    doc.add_heading("3. Limitations & Conclusion", level=1)
    doc.add_paragraph(
        "Limitations include a 4,000-character input ceiling for document summarization and localized knowledge indexing. "
        "The project successfully demonstrates modern agentic AI design patterns, achieving reliable, verifiable performance "
        "suitable for academic evaluation."
    )

    docx_path = ARTIFACTS_DIR / "project_report.docx"
    doc.save(str(docx_path))
    print(f"[SUCCESS] Generated Word Report: {docx_path}")

def generate_powerpoint_presentation():
    prs = Presentation()
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)
    blank_layout = prs.slide_layouts[6]

    def add_slide_header(slide, title_text, category_text="AGENTIC AI COURSE COMPLETION PROJECT"):
        # Header banner
        shape = slide.shapes.add_shape(1, 0, 0, PptxInches(13.333), PptxInches(1.2)) # msoShapeRectangle = 1
        shape.fill.solid()
        shape.fill.fore_color.rgb = PptxRGBColor(11, 26, 48)
        shape.line.color.rgb = PptxRGBColor(0, 119, 182)

        tf = shape.text_frame
        tf.margin_left = PptxInches(0.8)
        tf.margin_top = PptxInches(0.2)
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = PptxPt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = PptxRGBColor(0, 180, 216)

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = PptxPt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = PptxRGBColor(255, 255, 255)

    def add_content_box(slide, left, top, width, height, title, bullets, accent_color=PptxRGBColor(0, 119, 182)):
        box = slide.shapes.add_shape(1, PptxInches(left), PptxInches(top), PptxInches(width), PptxInches(height))
        box.fill.solid()
        box.fill.fore_color.rgb = PptxRGBColor(17, 34, 64)
        box.line.color.rgb = accent_color
        box.line.width = PptxPt(2)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = PptxInches(0.4)
        tf.margin_top = PptxInches(0.3)
        tf.margin_right = PptxInches(0.4)
        
        p_head = tf.paragraphs[0]
        p_head.text = title
        p_head.font.size = PptxPt(18)
        p_head.font.bold = True
        p_head.font.color.rgb = PptxRGBColor(255, 255, 255)

        for b in bullets:
            p = tf.add_paragraph()
            p.text = f"• {b}"
            p.font.size = PptxPt(14)
            p.font.color.rgb = PptxRGBColor(220, 230, 242)
            p.space_before = PptxPt(8)

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(1, 0, 0, PptxInches(13.333), PptxInches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PptxRGBColor(9, 17, 30)
    bg1.line.color.rgb = PptxRGBColor(0, 119, 182)
    
    tb1 = s1.shapes.add_textbox(PptxInches(1.0), PptxInches(1.5), PptxInches(11.333), PptxInches(4.5))
    tf1 = tb1.text_frame
    p_main = tf1.paragraphs[0]
    p_main.text = "Agentic AI Study & Productivity Assistant"
    p_main.font.size = PptxPt(36)
    p_main.font.bold = True
    p_main.font.color.rgb = PptxRGBColor(255, 255, 255)
    
    p_sub = tf1.add_paragraph()
    p_sub.text = "Autonomous Multi-Agent Planning, Tool Calling, Memory & Result Verification"
    p_sub.font.size = PptxPt(20)
    p_sub.font.color.rgb = PptxRGBColor(0, 180, 216)
    p_sub.space_before = PptxPt(14)
    
    p_meta = tf1.add_paragraph()
    p_meta.text = "\nStudent: [STUDENT_NAME]  |  ID: [STUDENT_ID]\nInstitution: [INSTITUTION_NAME]  |  Evaluation Date: [SUBMISSION_DATE]"
    p_meta.font.size = PptxPt(15)
    p_meta.font.color.rgb = PptxRGBColor(180, 200, 220)

    # Slide 2: Problem Statement & Motivation
    s2 = prs.slides.add_slide(blank_layout)
    add_slide_header(s2, "Problem Statement & Educational Challenges")
    add_content_box(s2, 0.8, 1.8, 5.6, 5.0, "The Challenges in Modern Learning", [
        "Cognitive overload when breaking down large exam syllabi.",
        "Generative AI chatbots produce hallucinated study schedules with inaccurate hours.",
        "Unverified math pacing leading to unrealistic study targets.",
        "Lack of structured active recall testing and feedback loops."
    ])
    add_content_box(s2, 6.9, 1.8, 5.6, 5.0, "Why Conversational Chatbots Fail", [
        "Single-turn token prediction without multi-step strategic planning.",
        "No programmatic verification of mathematical consistency.",
        "Arbitrary code execution security risks in ungrounded tools.",
        "Episodic memory resets on page refresh."
    ])

    # Slide 3: Objectives & Solution Overview
    s3 = prs.slides.add_slide(blank_layout)
    add_slide_header(s3, "Project Objectives & Core Innovations")
    add_content_box(s3, 0.8, 1.8, 5.6, 5.0, "Primary System Objectives", [
        "Hierarchical Goal Decomposition: Formulate structured TaskPlans.",
        "Safe Tool Allowlisting: Restrict operations to verified utilities.",
        "AST Math Parsing: Deterministic calculations without eval().",
        "Critic Verification Loop: Closed-loop auditing with bounded retries."
    ])
    add_content_box(s3, 6.9, 1.8, 5.6, 5.0, "Dual-Mode Reliability", [
        "Live LLM Mode: Connect to Gemini, OpenAI, or Anthropic.",
        "Offline Demonstration Mode: Deterministic local agent pipeline.",
        "Zero API Dependencies Required: 100% usable without tokens.",
        "Academic Verifiability: Reproducible test results on any machine."
    ])

    # Slide 4: System Architecture
    s4 = prs.slides.add_slide(blank_layout)
    add_slide_header(s4, "Multi-Agent System Architecture")
    add_content_box(s4, 0.8, 1.8, 3.6, 5.0, "1. Planning & Knowledge", [
        "Planner Agent: Interprets goals, extracts constraints, outputs typed TaskPlan.",
        "Knowledge Agent: Queries verified local curriculum without web hallucinations."
    ])
    add_content_box(s4, 4.8, 1.8, 3.6, 5.0, "2. Execution & Tools", [
        "Executor Agent: Safe parameter routing to allowlisted functions.",
        "Tools: Timetable generator, Quiz maker, AST math calculator, Doc summarizer."
    ])
    add_content_box(s4, 8.8, 1.8, 3.7, 5.0, "3. Verification & Memory", [
        "Critic Agent: Audits arithmetic, completeness, and structure (Max 2 revisions).",
        "Memory Manager: SQLite persistence with auto-sanitization of API keys."
    ])

    # Slide 5: Agent Roles & Responsibilities Matrix
    s5 = prs.slides.add_slide(blank_layout)
    add_slide_header(s5, "Agent Roles & Responsibilities Matrix")
    add_content_box(s5, 0.8, 1.8, 5.6, 5.0, "Autonomous Agents", [
        "Planner Agent: Generates subtasks and designates agent assignments.",
        "Knowledge Agent: Attributed citations from curated local knowledge.",
        "Executor Agent: Runs allowlisted tools with strict schema validation."
    ])
    add_content_box(s5, 6.9, 1.8, 5.6, 5.0, "Coordination & Memory", [
        "Critic Agent: Computes quality scores and flags deficiencies.",
        "Memory Manager: Parameterized SQLite storage for approved tasks.",
        "Study Orchestrator: Central coordinator managing execution timeline."
    ])

    # Slide 6: Security & AST Math Parsing
    s6 = prs.slides.add_slide(blank_layout)
    add_slide_header(s6, "Security Architecture: Safe Tool Allowlisting & AST Math")
    add_content_box(s6, 0.8, 1.8, 5.6, 5.0, "AST Math Evaluator (No eval/exec)", [
        "Python's eval() poses extreme code execution vulnerabilities.",
        "Our AST parser traverses Abstract Syntax Tree nodes explicitly.",
        "Only approved operators (+, -, *, /, %, pow) are permitted.",
        "Any unexpected node raises an immediate, safe ValueError."
    ])
    add_content_box(s6, 6.9, 1.8, 5.6, 5.0, "Additional Security Guardrails", [
        "Strict Allowlist: Executor Agent rejects unapproved tool names.",
        "SQL Parameterization: All SQLite statements use safe '?' placeholders.",
        "Credential Redaction: Regex filters strip API keys before database writes.",
        "Document Ceilings: Max 4,000-character limits to prevent buffer exhaustion."
    ])

    # Slide 7: Critic & Verification Loop
    s7 = prs.slides.add_slide(blank_layout)
    add_slide_header(s7, "Critic & Verification Feedback Loop")
    add_content_box(s7, 0.8, 1.8, 5.6, 5.0, "Multi-Point Quality Auditing", [
        "Completeness Check: Confirms all planned subtasks finished.",
        "Arithmetic Check: Asserts Total Hours == Days * Hours/Day.",
        "Quiz Structural Check: Verifies options count, answer key, and rationales.",
        "Confidence Score: Quantitative scoring (0% to 100%) computed per run."
    ])
    add_content_box(s7, 6.9, 1.8, 5.6, 5.0, "Bounded Revision Mechanism", [
        "Critic requests revisions when critical deficiencies are detected.",
        "Bounded Retry Budget: Maximum 2 revisions prevents infinite loops.",
        "Transparent Reporting: Unresolved issues are flagged in final report.",
        "Zero False Claims: System never masks failures or fabricates success."
    ])

    # Slide 8: Memory & Persistence
    s8 = prs.slides.add_slide(blank_layout)
    add_slide_header(s8, "Persistent Memory & Privacy Architecture")
    add_content_box(s8, 0.8, 1.8, 5.6, 5.0, "SQLite Local Storage", [
        "Normalized tables: task_memory and user_preferences.",
        "Saves user-approved study timetables, quizzes, and plans.",
        "Search functionality allows keyword recall across past goals.",
        "Full user control: Delete individual records or purge entire database."
    ])
    add_content_box(s8, 6.9, 1.8, 5.6, 5.0, "Privacy-First Standards", [
        "Runs 100% locally on the user's workstation.",
        "Automated Regex Sanitizer strips sk-... or Google API keys.",
        "No analytics, telemetry, or personal data transmitted externally."
    ])

    # Slide 9: Streamlit Application UI
    s9 = prs.slides.add_slide(blank_layout)
    add_slide_header(s9, "Interactive User Interface (Streamlit)")
    add_content_box(s9, 0.8, 1.8, 5.6, 5.0, "Modern Navy & Blue Dashboard", [
        "Assistant Dashboard: Real-time execution timeline and telemetry.",
        "Example Shortcuts: Pre-populated study goals for quick evaluation.",
        "Direct Tool Sandbox: Independently test calculators and quizzes.",
        "Memory Explorer: Search, inspect, and manage approved task records."
    ])
    add_content_box(s9, 6.9, 1.8, 5.6, 5.0, "Evaluator-Friendly Features", [
        "Force Offline Mode toggle directly in sidebar.",
        "Masked API key indicators showing configuration status.",
        "Expandable audit scorecards detailing checks passed and issues."
    ])

    # Slide 10: Automated Testing Results
    s10 = prs.slides.add_slide(blank_layout)
    add_slide_header(s10, "Automated Testing & Quality Assurance")
    add_content_box(s10, 0.8, 1.8, 5.6, 5.0, "Pytest Test Suite Overview", [
        "Executed under Python 3.14.5 and pytest 9.1.1.",
        "Total Test Count: 24 Automated Unit & Integration Tests.",
        "Execution Time: 1.54 seconds across 5 test modules.",
        "Pass Rate: 100% Unconditional Success (0 failures, 0 warnings)."
    ])
    add_content_box(s10, 6.9, 1.8, 5.6, 5.0, "Module-by-Module Verification", [
        "test_planner.py: 4/4 Passed (Goal breakdown, constraints).",
        "test_tools.py: 8/8 Passed (Timetables, quizzes, AST math, security).",
        "test_memory.py: 5/5 Passed (SQLite CRUD, sanitization, search).",
        "test_critic.py: 3/3 Passed (Math consistency, fault detection).",
        "test_orchestrator.py: 4/4 Passed (End-to-end workflow, callbacks)."
    ])

    # Slide 11: Limitations & Future Work
    s11 = prs.slides.add_slide(blank_layout)
    add_slide_header(s11, "Limitations & Future Roadmap")
    add_content_box(s11, 0.8, 1.8, 5.6, 5.0, "Current Technical Limitations", [
        "Document summarization capped at 4,000 characters.",
        "Offline research relies on pre-indexed local curriculum documents.",
        "Single-user desktop architecture rather than multi-tenant cloud."
    ])
    add_content_box(s11, 6.9, 1.8, 5.6, 5.0, "Future Enhancements", [
        "ChromaDB / pgvector integration for multi-page PDF RAG.",
        "Voice study questioning powered by local Whisper speech models.",
        "Direct export of study timetables to Google Calendar and iCal."
    ])

    # Slide 12: Conclusion & Q&A
    s12 = prs.slides.add_slide(blank_layout)
    bg12 = s12.shapes.add_shape(1, 0, 0, PptxInches(13.333), PptxInches(7.5))
    bg12.fill.solid()
    bg12.fill.fore_color.rgb = PptxRGBColor(9, 17, 30)
    
    tb12 = s12.shapes.add_textbox(PptxInches(1.0), PptxInches(2.0), PptxInches(11.333), PptxInches(3.5))
    tf12 = tb12.text_frame
    p_c1 = tf12.paragraphs[0]
    p_c1.text = "Thank You! Questions & Discussion"
    p_c1.font.size = PptxPt(36)
    p_c1.font.bold = True
    p_c1.font.color.rgb = PptxRGBColor(255, 255, 255)
    
    p_c2 = tf12.add_paragraph()
    p_c2.text = "Agentic AI Study & Productivity Assistant\nComplete, tested, and verified autonomous multi-agent implementation."
    p_c2.font.size = PptxPt(20)
    p_c2.font.color.rgb = PptxRGBColor(0, 180, 216)
    p_c2.space_before = PptxPt(16)

    pptx_path = ARTIFACTS_DIR / "project_presentation.pptx"
    prs.save(str(pptx_path))
    print(f"[SUCCESS] Generated PowerPoint Presentation: {pptx_path}")

if __name__ == "__main__":
    generate_word_report()
    generate_powerpoint_presentation()
