from pathlib import Path
from docx import Document
from pptx import Presentation

BASE = Path(__file__).resolve().parent.parent
docx_path = BASE / "artifacts" / "project_report.docx"
pptx_path = BASE / "artifacts" / "project_presentation.pptx"

print("--- WORD REPORT AUDIT ---")
if not docx_path.exists():
    print("ERROR: Word report does not exist!")
else:
    doc = Document(str(docx_path))
    print(f"File exists: {docx_path}")
    print(f"Total Paragraphs: {len(doc.paragraphs)}")
    print(f"Total Tables: {len(doc.tables)}")
    
    # Check for placeholders
    placeholders_found = set()
    for p in doc.paragraphs:
        txt = p.text
        for word in ["[STUDENT_NAME]", "[STUDENT_ID]", "[COURSE_OR_DEPARTMENT]", "[INSTITUTION_NAME]", "[SUBMISSION_DATE]", "[MENTOR_NAME]"]:
            if word in txt:
                placeholders_found.add(word)
    for t in doc.tables:
        for row in t.rows:
            for c in row.cells:
                for word in ["[STUDENT_NAME]", "[STUDENT_ID]", "[COURSE_OR_DEPARTMENT]", "[INSTITUTION_NAME]", "[SUBMISSION_DATE]", "[MENTOR_NAME]"]:
                    if word in c.text:
                        placeholders_found.add(word)
    print(f"Expected Student/Institution Placeholders verified: {sorted(list(placeholders_found))}")

print("\n--- POWERPOINT AUDIT ---")
if not pptx_path.exists():
    print("ERROR: PowerPoint does not exist!")
else:
    prs = Presentation(str(pptx_path))
    print(f"File exists: {pptx_path}")
    print(f"Total Slides: {len(prs.slides)}")
    for i, slide in enumerate(prs.slides):
        title = "No Title"
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines = shape.text_frame.text.strip().splitlines()
                title = lines[0] if len(lines) == 1 else lines[-1]
                break
        print(f"  Slide {i+1}: {title}")

print("\n[SUCCESS] Inspection script complete.")
